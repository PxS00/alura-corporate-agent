from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

from src.document_loader import DocumentLoader


def test_load_directory_extracts_supported_text_formats(tmp_path: Path):
    rh = tmp_path / "rh"
    financeiro = tmp_path / "financeiro"
    legal = tmp_path / "legal"
    operacional = tmp_path / "operacional"
    for directory in (rh, financeiro, legal, operacional):
        directory.mkdir()

    (rh / "politica.md").write_text("Política de férias vigente.", encoding="utf-8")
    (financeiro / "reembolso.csv").write_text(
        "tipo,limite\nAlimentação,R$ 100\n",
        encoding="utf-8",
    )
    (legal / "conduta.json").write_text(
        '{"regra": "Proteger dados pessoais"}',
        encoding="utf-8",
    )
    (operacional / "guia.html").write_text(
        "<html><body><main>Ative o MFA no primeiro dia.</main></body></html>",
        encoding="utf-8",
    )

    sections = DocumentLoader(tmp_path).load_directory()

    assert len(sections) == 4
    assert {section.metadata["category"] for section in sections} == {
        "rh",
        "financeiro",
        "legal",
        "operacional",
    }
    assert all(section.metadata["owner"] for section in sections)
    assert all(section.metadata["source_path"] for section in sections)


def test_load_pdf_extracts_text_and_page_metadata(tmp_path: Path):
    category = tmp_path / "rh"
    category.mkdir()
    path = category / "politica_ferias.pdf"

    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 800, "Politica de ferias corporativa em PDF.")
    pdf.showPage()
    pdf.save()

    sections = DocumentLoader(tmp_path).load_file(path)

    assert len(sections) == 1
    assert "Politica de ferias corporativa em PDF." in sections[0].text
    assert sections[0].metadata["page"] == 1
    assert sections[0].metadata["category"] == "rh"
    assert sections[0].metadata["file_type"] == "pdf"
    assert sections[0].metadata["source"] == "politica_ferias.pdf"


def test_load_docx_extracts_paragraphs_and_metadata(tmp_path: Path):
    category = tmp_path / "rh"
    category.mkdir()
    path = category / "beneficios.docx"

    document = DocxDocument()
    document.add_heading("Beneficios", level=1)
    document.add_paragraph("Auxilio educacao disponivel para colaboradores.")
    document.save(path)

    sections = DocumentLoader(tmp_path).load_file(path)

    assert len(sections) == 1
    assert "Beneficios" in sections[0].text
    assert "Auxilio educacao disponivel para colaboradores." in sections[0].text
    assert sections[0].metadata["section"] == "document"
    assert sections[0].metadata["category"] == "rh"
    assert sections[0].metadata["file_type"] == "docx"
    assert sections[0].metadata["source"] == "beneficios.docx"


def test_load_xlsx_extracts_structured_rows_and_sheet_metadata(tmp_path: Path):
    category = tmp_path / "financeiro"
    category.mkdir()
    path = category / "limites_reembolso.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Reembolsos"
    worksheet.append(["tipo", "limite"])
    worksheet.append(["Alimentacao", "R$ 120"])
    workbook.save(path)
    workbook.close()

    sections = DocumentLoader(tmp_path).load_file(path)

    assert len(sections) == 1
    assert "tipo: Alimentacao" in sections[0].text
    assert "limite: R$ 120" in sections[0].text
    assert sections[0].metadata["sheet"] == "Reembolsos"
    assert sections[0].metadata["category"] == "financeiro"
    assert sections[0].metadata["file_type"] == "xlsx"
    assert sections[0].metadata["source"] == "limites_reembolso.xlsx"


def test_load_pptx_extracts_slide_text_and_metadata(tmp_path: Path):
    category = tmp_path / "operacional"
    category.mkdir()
    path = category / "onboarding.pptx"

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Onboarding"
    slide.placeholders[1].text = "Ative o MFA no primeiro dia."
    presentation.save(path)

    sections = DocumentLoader(tmp_path).load_file(path)

    assert len(sections) == 1
    assert "Onboarding" in sections[0].text
    assert "Ative o MFA no primeiro dia." in sections[0].text
    assert sections[0].metadata["slide"] == 1
    assert sections[0].metadata["category"] == "operacional"
    assert sections[0].metadata["file_type"] == "pptx"
    assert sections[0].metadata["source"] == "onboarding.pptx"


def test_load_directory_removes_duplicate_content(tmp_path: Path):
    category = tmp_path / "rh"
    category.mkdir()
    content = "Mesmo documento corporativo."
    (category / "a.md").write_text(content, encoding="utf-8")
    (category / "b.md").write_text(content, encoding="utf-8")

    sections = DocumentLoader(tmp_path).load_directory()

    assert len(sections) == 1
