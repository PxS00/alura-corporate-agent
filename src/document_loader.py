from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from src.models import DocumentSection

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".md", ".csv", ".json", ".html", ".htm"}

CATEGORY_OWNERS = {
    "rh": "Recursos Humanos",
    "financeiro": "Financeiro",
    "operacional": "Operações",
    "legal": "Jurídico e Compliance",
    "comunicacao": "Comunicação Interna",
}


class DocumentLoader:
    """Extracts textual sections and metadata from supported corporate files."""

    def __init__(self, root: Path) -> None:
        self.root = root
        self._loaders: dict[str, Callable[[Path], list[DocumentSection]]] = {
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
            ".xlsx": self._load_xlsx,
            ".pptx": self._load_pptx,
            ".md": self._load_text,
            ".csv": self._load_csv,
            ".json": self._load_json,
            ".html": self._load_html,
            ".htm": self._load_html,
        }

    def load_directory(self) -> list[DocumentSection]:
        if not self.root.exists():
            raise FileNotFoundError(f"Diretório de documentos não encontrado: {self.root}")

        sections: list[DocumentSection] = []
        seen_hashes: set[str] = set()

        for path in sorted(self.root.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue

            for section in self.load_file(path):
                fingerprint = hashlib.sha256(section.text.encode("utf-8")).hexdigest()
                if fingerprint in seen_hashes:
                    continue
                seen_hashes.add(fingerprint)
                sections.append(section)

        return sections

    def load_file(self, path: Path) -> list[DocumentSection]:
        suffix = path.suffix.lower()
        loader = self._loaders.get(suffix)
        if loader is None:
            raise ValueError(f"Formato não suportado: {suffix}")

        sections = loader(path)
        common_metadata = self._common_metadata(path)

        return [
            DocumentSection(text=section.text, metadata={**common_metadata, **section.metadata})
            for section in sections
            if section.text.strip()
        ]

    def _common_metadata(self, path: Path) -> dict[str, str]:
        try:
            relative = path.relative_to(self.root)
        except ValueError:
            relative = Path(path.name)

        category = relative.parts[0].lower() if len(relative.parts) > 1 else "geral"
        updated_at = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat()

        return {
            "source": path.name,
            "source_path": relative.as_posix(),
            "category": category,
            "owner": CATEGORY_OWNERS.get(category, "Base de Conhecimento"),
            "file_type": path.suffix.lower().lstrip("."),
            "updated_at": updated_at,
        }

    @staticmethod
    def _load_pdf(path: Path) -> list[DocumentSection]:
        reader = PdfReader(str(path))
        return [
            DocumentSection(text=page.extract_text() or "", metadata={"page": index})
            for index, page in enumerate(reader.pages, start=1)
        ]

    @staticmethod
    def _load_docx(path: Path) -> list[DocumentSection]:
        document = Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        return [DocumentSection(text="\n\n".join(paragraphs), metadata={"section": "document"})]

    @staticmethod
    def _load_xlsx(path: Path) -> list[DocumentSection]:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
        sections: list[DocumentSection] = []
        try:
            for worksheet in workbook.worksheets:
                rows = list(worksheet.iter_rows(values_only=True))
                if not rows:
                    continue

                headers = [str(value).strip() if value is not None else f"coluna_{index + 1}" for index, value in enumerate(rows[0])]
                lines: list[str] = []
                for row in rows[1:]:
                    pairs = [
                        f"{headers[index]}: {value}"
                        for index, value in enumerate(row)
                        if value is not None and str(value).strip()
                    ]
                    if pairs:
                        lines.append(" | ".join(pairs))

                if lines:
                    sections.append(DocumentSection(text="\n".join(lines), metadata={"sheet": worksheet.title}))
        finally:
            workbook.close()
        return sections

    @staticmethod
    def _load_pptx(path: Path) -> list[DocumentSection]:
        presentation = Presentation(str(path))
        sections: list[DocumentSection] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if texts:
                sections.append(DocumentSection(text="\n".join(texts), metadata={"slide": slide_number}))
        return sections

    @staticmethod
    def _load_text(path: Path) -> list[DocumentSection]:
        return [DocumentSection(text=path.read_text(encoding="utf-8"), metadata={"section": "document"})]

    @staticmethod
    def _load_csv(path: Path) -> list[DocumentSection]:
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file)
            lines = [
                " | ".join(f"{key}: {value}" for key, value in row.items() if value and value.strip())
                for row in reader
            ]
        return [DocumentSection(text="\n".join(filter(None, lines)), metadata={"section": "table"})]

    @staticmethod
    def _load_json(path: Path) -> list[DocumentSection]:
        data = json.loads(path.read_text(encoding="utf-8"))
        return [
            DocumentSection(
                text=json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True),
                metadata={"section": "document"},
            )
        ]

    @staticmethod
    def _load_html(path: Path) -> list[DocumentSection]:
        soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
        for element in soup(["script", "style", "nav", "footer"]):
            element.decompose()
        text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
        return [DocumentSection(text=text, metadata={"section": "document"})]
